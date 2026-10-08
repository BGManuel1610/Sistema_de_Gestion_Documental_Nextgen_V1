import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_radicados_excel(radicados: list) -> io.BytesIO:
    """
    Genera un archivo .xlsx estandarizado con la información de radicados filtrada.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Radicados"

    # Encabezados
    headers = [
        "Código Radicado",
        "Remitente",
        "Asunto",
        "Tipo Documento",
        "Sede Destino",
        "Sede Actual",
        "Estado",
        "Fecha Creación",
        "Fecha Límite"
    ]
    
    ws.append(headers)

    # Estilos para encabezados
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Filas de datos
    for r in radicados:
        row = [
            r.get("rad_codigo", ""),
            r.get("rad_remitente", ""),
            r.get("rad_asunto", ""),
            r.get("rad_tipo_doc", ""),
            r.get("rad_sede_destino", ""),
            r.get("rad_sede_actual", ""),
            r.get("rad_estado", ""),
            str(r.get("rad_fecha_creacion", ""))[:19],
            str(r.get("rad_fecha_limite", ""))[:19]
        ]
        ws.append(row)

    # Ajustar ancho de columnas y bordes
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            if cell.row > 1:
                cell.font = Font(name="Calibri", size=10)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
